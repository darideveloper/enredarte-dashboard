"""Tests for the financial ledger (`finance` app)."""

import importlib
from decimal import Decimal

from django.apps import apps as global_apps
from django.contrib.admin.sites import site
from django.contrib.auth.models import User
from django.db.models import Sum
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from artworks.models import (
    Artist,
    Artwork,
    ArtworkOrder,
    ArtworkOrderCurrency,
    ArtworkOrderStatus,
    ArtworkStatus,
)
from artworks.services import apply_paid_transition
from finance import services as finance_services
from finance.admin import FinancialEntryAdmin
from finance.models import FinancialEntry
from subscriptions.models import ArtistSubscription, BillingPlan


class LedgerMixin:
    def make_artist(self, slug="frida", commission=25):
        return Artist.objects.create(
            name=f"Frida {slug}", slug=slug, email=f"{slug}@example.com", commission=commission
        )

    def make_order(self, artist, amount="1000.00", currency=ArtworkOrderCurrency.MXN, status=ArtworkStatus.RESERVED):
        artwork = Artwork.objects.create(
            artist=artist,
            year=2024,
            dimensions="10x10",
            price_mxn=amount,
            price_usd="100.00",
            status=status,
            slug=f"obra-{artist.slug}",
        )
        return ArtworkOrder.objects.create(
            artwork=artwork,
            currency=currency,
            amount=amount,
            buyer_email="buyer@example.com",
            status=ArtworkOrderStatus.PENDING_PAYMENT,
        )


class FinancialEntryModelTest(TestCase):
    def test_signed_convention_and_defaults(self):
        entry = FinancialEntry.objects.create(
            occurred_on="2026-01-15",
            kind=FinancialEntry.Kind.ARTIST_COMMISSION,
            amount=Decimal("-250.00"),
            currency=FinancialEntry.Currency.MXN,
        )
        self.assertEqual(entry.amount, Decimal("-250.00"))
        self.assertFalse(entry.reconciled)
        self.assertIn("Comisión de artista", str(entry))
        self.assertIn("-250", str(entry))

    def test_spanish_metadata(self):
        self.assertEqual(FinancialEntry._meta.verbose_name, "Movimiento")
        self.assertEqual(FinancialEntry._meta.verbose_name_plural, "Movimientos")
        self.assertEqual(FinancialEntry._meta.get_field("reconciled").verbose_name, "Conciliado")
        self.assertEqual(
            FinancialEntry.Kind.ARTWORK_SALE.label, "Venta de obra"
        )
        self.assertEqual(
            FinancialEntry.PaymentMethod.STRIPE.label, "En línea"
        )
        self.assertEqual(FinancialEntry.PaymentMethod.CASH.label, "Efectivo")

    def test_ordering_newest_first(self):
        a = FinancialEntry.objects.create(
            occurred_on="2026-01-01", kind=FinancialEntry.Kind.ARTWORK_SALE,
            amount="10.00", currency=FinancialEntry.Currency.MXN,
        )
        b = FinancialEntry.objects.create(
            occurred_on="2026-03-01", kind=FinancialEntry.Kind.ARTWORK_SALE,
            amount="10.00", currency=FinancialEntry.Currency.MXN,
        )
        self.assertEqual(list(FinancialEntry.objects.all()), [b, a])

    def test_blank_references_do_not_collide(self):
        for _ in range(2):
            FinancialEntry.objects.create(
                occurred_on="2026-01-01", kind=FinancialEntry.Kind.SUBSCRIPTION_PAYMENT,
                amount="10.00", currency=FinancialEntry.Currency.MXN, reference="",
            )
        self.assertEqual(FinancialEntry.objects.count(), 2)


class ArtworkSaleLedgerTest(LedgerMixin, TestCase):
    def test_record_sale_freezes_commission(self):
        artist = self.make_artist(commission=25)
        order = self.make_order(artist, amount="1000.00")
        order.stripe_payment_intent_id = "pi_123"
        order.save(update_fields=["stripe_payment_intent_id"])

        finance_services.record_artwork_sale(order)

        sale = FinancialEntry.objects.get(kind=FinancialEntry.Kind.ARTWORK_SALE)
        commission = FinancialEntry.objects.get(kind=FinancialEntry.Kind.ARTIST_COMMISSION)
        self.assertEqual(sale.amount, Decimal("1000.00"))
        self.assertEqual(commission.amount, Decimal("-250.00"))
        self.assertEqual(sale.currency, "MXN")
        self.assertEqual(sale.payment_method, FinancialEntry.PaymentMethod.STRIPE)
        self.assertEqual(sale.order_id, order.pk)

        artist.commission = 90
        artist.save(update_fields=["commission"])
        commission.refresh_from_db()
        self.assertEqual(commission.amount, Decimal("-250.00"))

    def test_zero_commission_only_income(self):
        artist = self.make_artist(commission=0)
        order = self.make_order(artist)
        finance_services.record_artwork_sale(order)
        self.assertEqual(FinancialEntry.objects.count(), 1)
        self.assertEqual(
            FinancialEntry.objects.get().kind, FinancialEntry.Kind.ARTWORK_SALE
        )

    def test_paid_transition_hook_and_idempotency(self):
        artist = self.make_artist(commission=20)
        order = self.make_order(artist)
        apply_paid_transition(order, payment_intent_id="pi_abc", buyer_email="b@e.com")
        self.assertEqual(FinancialEntry.objects.count(), 2)
        # Second transition is a no-op (status already paid) → no duplicate.
        apply_paid_transition(order, payment_intent_id="pi_abc")
        self.assertEqual(FinancialEntry.objects.count(), 2)


class RefundLedgerTest(LedgerMixin, TestCase):
    def test_refund_nets_to_zero_and_is_idempotent(self):
        artist = self.make_artist(commission=25)
        order = self.make_order(artist, amount="1000.00")
        apply_paid_transition(order, payment_intent_id="pi_ref")
        order.status = ArtworkOrderStatus.REFUNDED
        order.save(update_fields=["status"])

        finance_services.record_artwork_refund(order)
        finance_services.record_artwork_refund(order)

        entries = FinancialEntry.objects.filter(order=order)
        self.assertEqual(entries.count(), 4)
        self.assertEqual(sum((e.amount for e in entries), Decimal("0")), Decimal("0"))

    def test_refund_without_recorded_sale_is_noop(self):
        artist = self.make_artist(commission=25)
        order = self.make_order(artist, amount="1000.00")
        order.status = ArtworkOrderStatus.REFUNDED
        order.save(update_fields=["status"])

        finance_services.record_artwork_refund(order)

        self.assertFalse(FinancialEntry.objects.filter(order=order).exists())


class LedgerSignedSumTest(LedgerMixin, TestCase):
    def test_combined_signed_sum(self):
        artist = self.make_artist(commission=30)
        order = self.make_order(artist, amount="1000.00")
        apply_paid_transition(order, payment_intent_id="pi_sum")
        finance_services.record_subscription_payment(
            artist=artist,
            amount="500.00",
            currency="MXN",
            payment_method=FinancialEntry.PaymentMethod.STRIPE,
            reference="sub-sum",
        )
        total = FinancialEntry.objects.aggregate(net=Sum("amount"))["net"]
        self.assertEqual(total, Decimal("1200.00"))


class SubscriptionLedgerTest(LedgerMixin, TestCase):
    def _online_sub(self, artist):
        return ArtistSubscription.objects.create(
            artist=artist,
            status=ArtistSubscription.Status.ACTIVE,
            payment_method=ArtistSubscription.PaymentMethod.ONLINE,
            stripe_subscription_id="sub_x",
            stripe_customer_id="cus_x",
        )

    def _invoice_event(self, amount_paid=50000, currency="mxn", invoice_id="in_1"):
        return {
            "data": {
                "object": {
                    "id": invoice_id,
                    "customer": "cus_x",
                    "subscription": "sub_x",
                    "amount_paid": amount_paid,
                    "currency": currency,
                    "created": 1700000000,
                    "status_transitions": {"paid_at": 1700000000},
                    "lines": {"data": []},
                }
            }
        }

    def test_invoice_payment_snapshots_amount_and_dedupes(self):
        from subscriptions.webhooks import _handle_invoice_payment_succeeded

        artist = self.make_artist()
        sub = self._online_sub(artist)
        _handle_invoice_payment_succeeded(self._invoice_event())
        _handle_invoice_payment_succeeded(self._invoice_event())

        entry = FinancialEntry.objects.get(kind=FinancialEntry.Kind.SUBSCRIPTION_PAYMENT)
        self.assertEqual(entry.amount, Decimal("500.00"))
        self.assertEqual(entry.currency, "MXN")
        self.assertEqual(entry.payment_method, FinancialEntry.PaymentMethod.STRIPE)
        self.assertEqual(entry.reference, "in_1")
        self.assertEqual(entry.subscription_id, sub.pk)
        self.assertEqual(FinancialEntry.objects.count(), 1)

    def test_stripe_handler_ignores_cash_rows(self):
        from subscriptions.webhooks import _handle_invoice_payment_succeeded

        artist = self.make_artist()
        ArtistSubscription.objects.create(
            artist=artist,
            status=ArtistSubscription.Status.ACTIVE,
            payment_method=ArtistSubscription.PaymentMethod.CASH,
            stripe_subscription_id="sub_x",
            stripe_customer_id="cus_x",
        )
        _handle_invoice_payment_succeeded(self._invoice_event())
        self.assertFalse(FinancialEntry.objects.exists())

    def test_cash_snapshot_and_plan_change(self):
        plan = BillingPlan.get_solo()
        plan.amount = Decimal("500.00")
        plan.currency = "MXN"
        plan.save()

        artist = self.make_artist()
        sub = ArtistSubscription.objects.create(
            artist=artist,
            status=ArtistSubscription.Status.PENDING,
            payment_method=ArtistSubscription.PaymentMethod.CASH,
        )
        sub.current_period_end = timezone.now()
        entry = finance_services.record_cash_subscription_payment(sub)
        self.assertEqual(entry.amount, Decimal("500.00"))
        self.assertEqual(entry.payment_method, FinancialEntry.PaymentMethod.CASH)

        plan.amount = Decimal("900.00")
        plan.save()
        entry.refresh_from_db()
        self.assertEqual(entry.amount, Decimal("500.00"))


class BackfillTest(LedgerMixin, TestCase):
    def setUp(self):
        self.mod = importlib.import_module(
            "finance.migrations.0002_backfill_artwork_sales"
        )

    def test_backfill_paid_only_and_idempotent(self):
        artist = self.make_artist(commission=25)
        paid = self.make_order(artist, amount="1000.00")
        paid.paid_at = "2026-01-10T10:00:00Z"
        paid.status = ArtworkOrderStatus.DATA_COMPLETE
        paid.stripe_payment_intent_id = "pi_back"
        paid.save(update_fields=["paid_at", "status", "stripe_payment_intent_id"])

        unpaid = self.make_order(
            self.make_artist(slug="otro", commission=10), status=ArtworkStatus.AVAILABLE
        )

        self.mod.backfill(global_apps, None)
        self.mod.backfill(global_apps, None)

        self.assertEqual(FinancialEntry.objects.filter(order=paid).count(), 2)
        self.assertFalse(FinancialEntry.objects.filter(order=unpaid).exists())

        self.mod.unbackfill(global_apps, None)
        self.assertFalse(FinancialEntry.objects.filter(order=paid).exists())


class FinanceAdminTest(LedgerMixin, TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            username="admin", email="a@a.com", password="password123"
        )
        self.client.force_login(self.user)
        self.url = reverse("admin:finance_financialentry_changelist")
        self.artist = self.make_artist(commission=25)
        self.sale = FinancialEntry.objects.create(
            occurred_on="2026-01-15", kind=FinancialEntry.Kind.ARTWORK_SALE,
            amount="1000.00", currency=FinancialEntry.Currency.MXN,
            artist=self.artist, reference="sale-1",
        )
        self.commission = FinancialEntry.objects.create(
            occurred_on="2026-01-15", kind=FinancialEntry.Kind.ARTIST_COMMISSION,
            amount="-250.00", currency=FinancialEntry.Currency.MXN,
            artist=self.artist, reference="sale-1-c",
        )
        self.admin = site._registry[FinancialEntry]

    def test_registered_with_expected_config(self):
        self.assertIsInstance(self.admin, FinancialEntryAdmin)
        self.assertEqual(self.admin.list_editable, ("reconciled",))
        self.assertFalse(self.admin.has_add_permission(None))
        self.assertFalse(self.admin.has_delete_permission(None))
        self.assertNotIn("reconciled", self.admin.readonly_fields)

    def test_inline_reconcile_from_list(self):
        data = {
            "form-TOTAL_FORMS": "1",
            "form-INITIAL_FORMS": "1",
            "form-MIN_NUM_FORMS": "0",
            "form-MAX_NUM_FORMS": "1000",
            "form-0-id": str(self.sale.pk),
            "form-0-reconciled": "on",
            "_save": "Guardar",
        }
        response = self.client.post(self.url, data)
        self.assertIn(response.status_code, (200, 302))
        self.sale.refresh_from_db()
        self.assertTrue(self.sale.reconciled)
        self.commission.refresh_from_db()
        self.assertFalse(self.commission.reconciled)

    def test_spanish_labels_and_totals(self):
        response = self.client.get(self.url)
        content = response.content.decode()
        self.assertIn("Conciliado", content)
        self.assertIn("Vista actual", content)
        self.assertIn("Total del sistema", content)

    def test_totals_filtered_vs_full(self):
        response = self.client.get(self.url, {"kind": "artist_commission"})
        totals = response.context["ledger_totals"]
        filtered = totals["filtered"][0]
        full = totals["full"][0]
        self.assertEqual(filtered["income"], Decimal("0"))
        self.assertEqual(filtered["expenses"], Decimal("-250.00"))
        self.assertEqual(full["income"], Decimal("1000.00"))
        self.assertEqual(full["net"], Decimal("750.00"))

    def test_full_total_shown_with_no_results(self):
        other = self.make_artist(slug="nadie", commission=0)
        response = self.client.get(self.url, {"artist__id__exact": str(other.pk)})
        totals = response.context["ledger_totals"]
        self.assertEqual(totals["filtered"], [])
        self.assertEqual(totals["full"][0]["net"], Decimal("750.00"))

    def test_month_drilldown_scopes_filtered_but_not_full(self):
        FinancialEntry.objects.create(
            occurred_on="2026-02-10", kind=FinancialEntry.Kind.ARTWORK_SALE,
            amount="500.00", currency=FinancialEntry.Currency.MXN,
            artist=self.artist, reference="sale-2",
        )
        response = self.client.get(
            self.url, {"occurred_on__year": "2026", "occurred_on__month": "2"}
        )
        totals = response.context["ledger_totals"]
        self.assertEqual(totals["filtered"][0]["count"], 1)
        self.assertEqual(totals["filtered"][0]["income"], Decimal("500.00"))
        self.assertEqual(totals["full"][0]["net"], Decimal("1250.00"))

    def test_totals_split_by_currency(self):
        FinancialEntry.objects.create(
            occurred_on="2026-01-20", kind=FinancialEntry.Kind.SUBSCRIPTION_PAYMENT,
            amount="100.00", currency=FinancialEntry.Currency.USD,
            artist=self.artist, reference="sub-usd",
        )
        response = self.client.get(self.url)
        currencies = {row["currency"] for row in response.context["ledger_totals"]["full"]}
        self.assertEqual(currencies, {"MXN", "USD"})

    def test_add_and_delete_blocked_over_http(self):
        add = self.client.get(reverse("admin:finance_financialentry_add"))
        delete = self.client.get(
            reverse("admin:finance_financialentry_delete", args=[self.sale.pk])
        )
        self.assertEqual(add.status_code, 403)
        self.assertEqual(delete.status_code, 403)

    def test_filters_payment_method_and_currency(self):
        cash = FinancialEntry.objects.create(
            occurred_on="2026-01-16", kind=FinancialEntry.Kind.SUBSCRIPTION_PAYMENT,
            amount="500.00", currency=FinancialEntry.Currency.MXN,
            payment_method=FinancialEntry.PaymentMethod.CASH,
            artist=self.artist, reference="cash-f",
        )
        usd = FinancialEntry.objects.create(
            occurred_on="2026-01-17", kind=FinancialEntry.Kind.SUBSCRIPTION_PAYMENT,
            amount="100.00", currency=FinancialEntry.Currency.USD,
            payment_method=FinancialEntry.PaymentMethod.STRIPE,
            artist=self.artist, reference="usd-f",
        )
        by_method = self.client.get(self.url, {"payment_method__exact": "cash"})
        ids = {e.pk for e in by_method.context["cl"].result_list}
        self.assertEqual(ids, {cash.pk})

        by_currency = self.client.get(self.url, {"currency__exact": "USD"})
        ids = {e.pk for e in by_currency.context["cl"].result_list}
        self.assertEqual(ids, {usd.pk})

    def test_filter_reconciled(self):
        self.sale.reconciled = True
        self.sale.save(update_fields=["reconciled"])
        response = self.client.get(self.url, {"reconciled__exact": "1"})
        ids = {e.pk for e in response.context["cl"].result_list}
        self.assertEqual(ids, {self.sale.pk})

    def test_amount_color_coding(self):
        content = self.client.get(self.url).content.decode()
        self.assertIn("#166534", content)
        self.assertIn("#991b1b", content)

    def test_internal_no_public_surface(self):
        import importlib.util

        self.assertIsNone(importlib.util.find_spec("finance.urls"))
        self.assertIsNone(importlib.util.find_spec("finance.serializers"))


class ExcelExportUnitTests(LedgerMixin, TestCase):
    def setUp(self):
        self.artist = self.make_artist(commission=20)
        self.order = self.make_order(self.artist, amount="1500.00")
        self.entry = FinancialEntry.objects.create(
            occurred_on="2026-03-15",
            kind=FinancialEntry.Kind.ARTWORK_SALE,
            amount=Decimal("1500.00"),
            currency=FinancialEntry.Currency.MXN,
            artist=self.artist,
            order=self.order,
            reference="sale-unit-1",
        )

    def test_sanitize_sheet_name(self):
        from utils.excel_export import sanitize_sheet_name

        self.assertEqual(sanitize_sheet_name("CleanName"), "CleanName")
        self.assertEqual(sanitize_sheet_name("Invalid:/*?[]Chars"), "Invalid______Chars")
        long_name = "A" * 40
        self.assertEqual(len(sanitize_sheet_name(long_name)), 31)

        existing = {"Test", "Test_2"}
        self.assertEqual(sanitize_sheet_name("Test", existing), "Test_3")

    def test_serialize_value(self):
        from utils.excel_export import serialize_value

        field_amount = FinancialEntry._meta.get_field("amount")
        self.assertEqual(serialize_value(field_amount, self.entry), 1500.0)

        field_occurred = FinancialEntry._meta.get_field("occurred_on")
        self.assertEqual(str(serialize_value(field_occurred, self.entry)), "2026-03-15")

        field_reconciled = FinancialEntry._meta.get_field("reconciled")
        self.assertEqual(serialize_value(field_reconciled, self.entry), False)

    def test_columns_for_model_splits_foreign_keys(self):
        from utils.excel_export import columns_for_model

        cols = columns_for_model(FinancialEntry)
        headers = [c[0] for c in cols]

        self.assertIn("artist__str__", headers)
        self.assertIn("artist_id", headers)
        self.assertIn("order__str__", headers)
        self.assertIn("order_id", headers)
        self.assertIn("subscription__str__", headers)
        self.assertIn("subscription_id", headers)

        getters = {c[0]: c[1] for c in cols}
        self.assertEqual(getters["artist__str__"](self.entry), str(self.artist))
        self.assertEqual(getters["artist_id"](self.entry), self.artist.pk)
        self.assertEqual(getters["order__str__"](self.entry), str(self.order))
        self.assertEqual(getters["order_id"](self.entry), self.order.pk)
        self.assertEqual(getters["subscription__str__"](self.entry), "")
        self.assertEqual(getters["subscription_id"](self.entry), "")

    def test_single_sheet_workbook_building(self):
        import openpyxl
        from io import BytesIO
        from utils.excel_export import build_workbook_for_queryset

        wb = build_workbook_for_queryset(
            FinancialEntry,
            FinancialEntry.objects.filter(pk=self.entry.pk),
            include_related=False,
        )
        self.assertEqual(len(wb.sheetnames), 1)
        self.assertEqual(wb.sheetnames[0], "Movimientos")

        ws = wb.active
        self.assertEqual(ws.freeze_panes, "A2")

        headers = [cell.value for cell in ws[1]]
        self.assertIn("artist__str__", headers)
        self.assertIn("artist_id", headers)

        row_vals = [cell.value for cell in ws[2]]
        artist_str_idx = headers.index("artist__str__")
        artist_id_idx = headers.index("artist_id")
        self.assertEqual(row_vals[artist_str_idx], str(self.artist))
        self.assertEqual(row_vals[artist_id_idx], self.artist.pk)


class ExcelExportRelatedDataTests(LedgerMixin, TestCase):
    def setUp(self):
        self.artist = self.make_artist(slug="diego", commission=30)
        self.order = self.make_order(self.artist, amount="2000.00")
        self.sale = FinancialEntry.objects.create(
            occurred_on="2026-02-01",
            kind=FinancialEntry.Kind.ARTWORK_SALE,
            amount=Decimal("2000.00"),
            currency=FinancialEntry.Currency.MXN,
            artist=self.artist,
            order=self.order,
            reference="sale-diego",
        )

    def test_multi_sheet_export_referenced_rows_only(self):
        from utils.excel_export import build_workbook_for_queryset

        unreferenced_artist = self.make_artist(slug="unref", commission=10)

        wb = build_workbook_for_queryset(
            FinancialEntry,
            FinancialEntry.objects.filter(pk=self.sale.pk),
            include_related=True,
        )

        sheetnames = wb.sheetnames
        self.assertIn("Movimientos", sheetnames)
        self.assertIn("Artistas", sheetnames)
        self.assertIn(str(ArtworkOrder._meta.verbose_name_plural), sheetnames)
        self.assertNotIn("Suscripciones de artistas", sheetnames)

        artist_sheet = wb["Artistas"]
        artist_headers = [c.value for c in artist_sheet[1]]
        name_idx = artist_headers.index("Nombre")
        artist_names = [artist_sheet.cell(row=r, column=name_idx + 1).value for r in range(2, artist_sheet.max_row + 1)]
        self.assertIn(self.artist.name, artist_names)
        self.assertNotIn(unreferenced_artist.name, artist_names)

        order_sheet = wb[str(ArtworkOrder._meta.verbose_name_plural)]
        self.assertEqual(order_sheet.freeze_panes, "A2")
        self.assertEqual(order_sheet.max_row, 2)


class ExcelExportAdminIntegrationTests(LedgerMixin, TestCase):
    def setUp(self):
        self.admin_user = User.objects.create_superuser(
            username="excel_admin", email="admin@example.com", password="password123"
        )
        self.client.force_login(self.admin_user)
        self.url = reverse("admin:finance_financialentry_changelist")
        self.artist = self.make_artist(slug="exportart", commission=20)
        self.entry = FinancialEntry.objects.create(
            occurred_on="2026-03-01",
            kind=FinancialEntry.Kind.ARTWORK_SALE,
            amount=Decimal("1200.00"),
            currency=FinancialEntry.Currency.MXN,
            artist=self.artist,
            reference="export-ref-1",
        )

    def test_export_selected_changelist_action(self):
        import openpyxl
        from io import BytesIO

        response = self.client.post(
            self.url,
            {
                "action": "export_selected",
                "_selected_action": [str(self.entry.pk)],
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment; filename=\"financialentry_", response["Content-Disposition"])

        wb = openpyxl.load_workbook(BytesIO(response.content))
        self.assertEqual(len(wb.sheetnames), 1)
        self.assertEqual(wb.sheetnames[0], "Movimientos")
        ws = wb.active
        self.assertEqual(ws.max_row, 2)

    def test_export_selected_with_related_changelist_action(self):
        import openpyxl
        from io import BytesIO

        response = self.client.post(
            self.url,
            {
                "action": "export_selected_with_related",
                "_selected_action": [str(self.entry.pk)],
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        wb = openpyxl.load_workbook(BytesIO(response.content))
        self.assertIn("Movimientos", wb.sheetnames)
        self.assertIn("Artistas", wb.sheetnames)

    def test_empty_selection_shows_warning(self):
        response = self.client.post(
            self.url,
            {
                "action": "export_selected",
            },
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        # Content-Type should remain HTML, not an Excel download
        self.assertIn("text/html", response["Content-Type"])

    def test_export_requires_view_permission(self):
        non_staff = User.objects.create_user(
            username="regular_user", email="reg@example.com", password="password123"
        )
        self.client.force_login(non_staff)
        response = self.client.post(
            self.url,
            {
                "action": "export_selected",
                "_selected_action": [str(self.entry.pk)],
            },
        )
        self.assertEqual(response.status_code, 302)

    def test_no_global_export_button_in_admin(self):
        response = self.client.get(self.url)
        content = response.content.decode()
        self.assertNotIn("Export all app data", content)
        self.assertNotIn("export_all", content)

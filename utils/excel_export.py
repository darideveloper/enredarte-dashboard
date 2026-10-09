"""Excel workbook export utilities for Django models using openpyxl.

Follows the patterns documented in docs/django-excel-export.md.
"""

import json
from datetime import date, datetime, time
from decimal import Decimal

from django.db import models
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FALLBACK_COLOR = "2563EB"  # Generic blue (blue-600)


def _primary_color():
    """Return the primary brand hex color without leading '#'.

    Defaults to FALLBACK_COLOR.
    """
    return FALLBACK_COLOR


def sanitize_sheet_name(name, existing=None):
    """Sanitize and clamp sheet names for Excel compatibility.

    Excel sheet names must not exceed 31 characters, cannot contain
    ':\\/?*[]', and must be unique within the workbook.
    """
    if not name:
        base = "Sheet"
    else:
        cleaned = str(name).strip("'")
        for ch in [":", "\\", "/", "?", "*", "[", "]"]:
            cleaned = cleaned.replace(ch, "_")
        cleaned = cleaned.strip()
        base = cleaned[:31] if cleaned else "Sheet"

    existing_set = set(existing) if existing else set()
    if base not in existing_set:
        return base

    counter = 2
    while True:
        suffix = f"_{counter}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        if candidate not in existing_set:
            return candidate
        counter += 1


def serialize_value(field, obj):
    """Serialize a concrete model field value into an Excel-friendly primitive."""
    attname = getattr(field, "attname", field.name)
    value = getattr(obj, attname, None)
    if value is None:
        return ""

    if isinstance(field, models.FileField):
        return getattr(value, "name", "") or ""

    if isinstance(field, models.JSONField):
        return json.dumps(value, ensure_ascii=False)

    if isinstance(field, models.BooleanField):
        return bool(value)

    if isinstance(field, models.DecimalField):
        return float(value)

    if isinstance(
        field,
        (
            models.IntegerField,
            models.AutoField,
            models.BigAutoField,
            models.SmallIntegerField,
            models.BigIntegerField,
            models.PositiveIntegerField,
            models.PositiveSmallIntegerField,
            models.PositiveBigIntegerField,
        ),
    ):
        return int(value)

    if isinstance(field, models.DateTimeField):
        if isinstance(value, datetime):
            return value.replace(tzinfo=None) if value.tzinfo else value
        return value

    if isinstance(field, (models.DateField, models.TimeField)):
        return value

    return str(value)


def columns_for_model(model):
    """Return a list of (header, getter, field) tuples for all concrete fields.

    Forward foreign keys (ForeignKey and OneToOneField) are rendered as two columns:
    1. {field.name}__str__: string representation of related object.
    2. {field.attname}: database ID of related object.
    """
    columns = []
    for field in model._meta.concrete_fields:
        if isinstance(field, (models.ForeignKey, models.OneToOneField)) and not field.auto_created:
            fk_name = field.name
            fk_attname = field.attname
            header_str = f"{fk_name}__str__"
            getter_str = (
                lambda obj, name=fk_name: str(getattr(obj, name))
                if getattr(obj, name, None) is not None
                else ""
            )
            columns.append((header_str, getter_str, field))

            header_id = fk_attname
            getter_id = (
                lambda obj, att=fk_attname: getattr(obj, att)
                if getattr(obj, att, None) is not None
                else ""
            )
            columns.append((header_id, getter_id, field))
        else:
            header = str(field.verbose_name or field.name)
            getter = lambda obj, f=field: serialize_value(f, obj)
            columns.append((header, getter, field))
    return columns


def get_related_targets(model):
    """Return a list of distinct forward ForeignKey/OneToOne target models."""
    targets = []
    seen = set()
    for field in model._meta.concrete_fields:
        if isinstance(field, (models.ForeignKey, models.OneToOneField)) and not field.auto_created:
            target = field.remote_field.model
            if target and target not in seen:
                seen.add(target)
                targets.append(target)
    return targets


def style_sheet(ws):
    """Apply styling to a worksheet: header fill, freeze A2, zebra rows, number formatting."""
    primary = _primary_color().lstrip("#")
    header_fill = PatternFill(
        start_color=f"FF{primary}",
        end_color=f"FF{primary}",
        fill_type="solid",
    )
    header_font = Font(color="FFFFFF", bold=True, size=11)
    header_alignment = Alignment(horizontal="center", vertical="center")
    bottom_border = Border(bottom=Side(style="thin", color="D1D5DB"))

    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = bottom_border

    ws.freeze_panes = "A2"

    even_fill = PatternFill(start_color="FFF9FAFB", end_color="FFF9FAFB", fill_type="solid")
    odd_fill = PatternFill(start_color="FFFFFFFF", end_color="FFFFFFFF", fill_type="solid")

    for row_idx in range(2, ws.max_row + 1):
        row_fill = even_fill if row_idx % 2 == 0 else odd_fill
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = row_fill
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"


def autosize_columns(ws):
    """Auto-size columns clamped between 10 and 50 characters."""
    for col in ws.columns:
        col_letter = col[0].column_letter
        max_len = 0
        for cell in col:
            if cell.value is not None:
                cell_len = len(str(cell.value))
                if cell_len > max_len:
                    max_len = cell_len
        width = max(10, min(50, int(max_len * 1.2) + 2))
        ws.column_dimensions[col_letter].width = width


def _write_sheet_for_model(wb, model, queryset, existing_names, is_first=False):
    """Write headers and rows for a model queryset to a worksheet."""
    title_raw = str(model._meta.verbose_name_plural or model._meta.model_name)
    title = sanitize_sheet_name(title_raw, existing_names)
    existing_names.add(title)

    if is_first and wb.active and len(wb.sheetnames) == 1 and wb.sheetnames[0] == "Sheet":
        ws = wb.active
        ws.title = title
    else:
        ws = wb.create_sheet(title=title)

    cols = columns_for_model(model)
    ws.append([header for header, _, _ in cols])

    fk_names = [
        f.name
        for f in model._meta.concrete_fields
        if isinstance(f, (models.ForeignKey, models.OneToOneField)) and not f.auto_created
    ]
    qs = queryset
    if fk_names and hasattr(qs, "select_related"):
        try:
            qs = qs.select_related(*fk_names)
        except Exception:
            pass

    iterator = qs.iterator() if hasattr(qs, "iterator") else iter(qs)
    for obj in iterator:
        ws.append([getter(obj) for _, getter, _ in cols])

    style_sheet(ws)
    autosize_columns(ws)
    return ws


def build_workbook_for_queryset(model, queryset, include_related=False, existing_workbook=None):
    """Build an Excel workbook for the given model and queryset.

    Args:
        model: Django model class.
        queryset: QuerySet of model instances to export.
        include_related: If True, adds separate worksheets for each referenced forward target model.
        existing_workbook: Optional openpyxl Workbook instance to add sheets to.
    """
    wb = existing_workbook or Workbook()
    existing_names = set(wb.sheetnames)

    _write_sheet_for_model(wb, model, queryset, existing_names, is_first=True)

    if include_related:
        fk_fields = [
            f
            for f in model._meta.concrete_fields
            if isinstance(f, (models.ForeignKey, models.OneToOneField)) and not f.auto_created
        ]
        target_to_ids = {}
        for f in fk_fields:
            target = f.remote_field.model
            if not target:
                continue
            if hasattr(queryset, "values_list"):
                ids = set(
                    queryset.exclude(**{f"{f.attname}__isnull": True}).values_list(
                        f.attname, flat=True
                    )
                )
            else:
                ids = {getattr(obj, f.attname) for obj in queryset if getattr(obj, f.attname) is not None}
            target_to_ids.setdefault(target, set()).update(ids)

        for target, ids in target_to_ids.items():
            valid_ids = [i for i in ids if i is not None]
            if not valid_ids:
                continue
            target_qs = target.objects.filter(pk__in=valid_ids)
            if not target_qs.exists():
                continue
            _write_sheet_for_model(wb, target, target_qs, existing_names, is_first=False)

    return wb
